import csv
import base64
import uuid
import logging
from datetime import datetime, time, timezone
from io import BytesIO, StringIO
from functools import lru_cache
from pathlib import Path

from fastapi import APIRouter, Depends, File, Form, HTTPException, Response, UploadFile, status
from fastapi.responses import FileResponse
from sqlalchemy import func, select, tuple_
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from backend.app.config import get_settings
from backend.app.database import get_db
from backend.app.models import (
    AttendanceRecord,
    Class,
    ClassEnrollment,
    ClassSession,
    ConsentStatus,
    FaceEnrollmentDraft,
    FaceTemplate,
    SessionStatus,
    Student,
    User,
    utc_now,
)
from backend.app.schemas import (
    AttendanceManualUpdate,
    AttendanceBoardRead,
    AttendanceBoardRow,
    AttendanceRead,
    AttendanceRequestCreate,
    AttendanceRequestRead,
    AttendanceStatsRead,
    ClassCreate,
    ClassRead,
    DashboardSummaryRead,
    FaceEnrollmentProgress,
    FaceEnrollmentScanRequest,
    FaceImportItem,
    FaceImportResponse,
    FaceTemplateRead,
    FrameScanRequest,
    FrameScanResponse,
    StreamScanRequest,
    StreamScanResponse,
    LoginRequest,
    RecognitionEventCreate,
    RecognitionEventRead,
    SessionCreate,
    SessionPolicyUpdate,
    SessionRead,
    SessionSummaryRead,
    StudentFaceSummary,
    StudentCreate,
    StudentRead,
    StudentImportResult,
    StudentUpdate,
    UserCreate,
    UserRead,
)
from backend.app.security import (
    bearer,
    allowed_class_ids,
    get_current_user,
    hash_password,
    issue_session,
    token_digest,
    require_admin,
    require_teacher_or_admin,
    require_user_or_admin,
    verify_camera_token,
    verify_password,
)
from backend.app.services.attendance import (
    close_session,
    create_recognition_event,
    open_session,
    update_attendance_manually,
)
from backend.app.services.face_templates import build_default_recognizer, register_face_bytes, register_face_image
from backend.app.services.audit import write_audit

router = APIRouter()
logger = logging.getLogger(__name__)


@lru_cache(maxsize=4)
def _camera_confirmations(min_hits: int, window_seconds: int):
    from ai_pipeline.temporal import CameraConfirmationStore

    return CameraConfirmationStore(min_hits, window_seconds)


def _localized_face_error(reason: str | None) -> str:
    text = reason or "Không thể phân tích khuôn mặt"
    translations = {
        "No face detected": "Không phát hiện khuôn mặt",
        "Multiple faces detected": "Chỉ cho phép một khuôn mặt trong khung hình",
        "Cannot decode image": "Không đọc được dữ liệu ảnh",
        "Empty face crop": "Vùng khuôn mặt không hợp lệ",
        "Face is too small": "Khuôn mặt quá nhỏ, vui lòng tiến gần camera",
        "Image is too dark": "Khung hình quá tối, vui lòng tăng ánh sáng",
        "Image is overexposed": "Khung hình quá sáng, vui lòng điều chỉnh ánh sáng",
        "Image is blurry": "Ảnh bị mờ, vui lòng giữ yên khuôn mặt",
        "Image has insufficient detail": "Ảnh thiếu chi tiết hoặc tương phản, vui lòng đổi vị trí và ánh sáng",
        "Face is clipped": "Khuôn mặt bị cắt ở mép ảnh, vui lòng đưa toàn bộ khuôn mặt vào khung hình",
    }
    for prefix, message in translations.items():
        if text.startswith(prefix):
            return message
    return text


def _attendance_evidence_url(attendance: AttendanceRecord | None) -> str | None:
    if attendance is None or attendance.evidence_image_path is None:
        return None
    return f"/attendance/{attendance.id}/evidence"


def _save_attendance_evidence(
    db: Session,
    *,
    attendance: AttendanceRecord | None,
    image_bytes: bytes,
    face_box=None,
) -> None:
    if attendance is None or attendance.evidence_image_path is not None:
        return
    from ai_pipeline.evidence import face_evidence

    evidence_bytes = face_evidence(image_bytes, face_box)

    storage_dir = get_settings().storage_dir / "attendance_evidence" / f"session-{attendance.session_id}"
    storage_dir.mkdir(parents=True, exist_ok=True)
    image_path = storage_dir / f"attendance-{attendance.id}-student-{attendance.student_id}.jpg"
    image_path.write_bytes(evidence_bytes)
    attendance.evidence_image_path = str(image_path)
    db.commit()
    db.refresh(attendance)


def _build_attendance_board(db: Session, class_session: ClassSession) -> AttendanceBoardRead:
    rows = db.execute(
        select(Student, AttendanceRecord)
        .join(ClassEnrollment, ClassEnrollment.student_id == Student.id)
        .outerjoin(
            AttendanceRecord,
            (AttendanceRecord.student_id == Student.id) & (AttendanceRecord.session_id == class_session.id),
        )
        .where(ClassEnrollment.class_id == class_session.class_id)
        .where(Student.deleted_at.is_(None))
        .order_by(Student.student_code)
    ).all()
    board_rows: list[AttendanceBoardRow] = []
    for student, attendance in rows:
        if attendance is None:
            board_rows.append(
                AttendanceBoardRow(
                    student_id=student.id,
                    student_code=student.student_code,
                    full_name=student.full_name,
                    class_label=student.class_label,
                    status="pending",
                    source=None,
                    recorded_at=None,
                    late_minutes=0,
                    alert="Chưa điểm danh",
                )
            )
            continue

        if attendance.status == "late":
            alert = f"Trễ {attendance.late_minutes} phút"
        elif attendance.status == "present":
            alert = "Đúng giờ"
        elif attendance.status == "absent":
            alert = "Vắng"
        else:
            alert = attendance.status
        board_rows.append(
            AttendanceBoardRow(
                student_id=student.id,
                student_code=student.student_code,
                full_name=student.full_name,
                class_label=student.class_label,
                status=attendance.status,
                source=attendance.source,
                recorded_at=attendance.recorded_at,
                late_minutes=attendance.late_minutes,
                alert=alert,
                evidence_image_url=_attendance_evidence_url(attendance),
            )
        )

    return AttendanceBoardRead(
        session_id=class_session.id,
        class_id=class_session.class_id,
        expected_start_time=class_session.expected_start_time,
        late_grace_minutes=class_session.late_grace_minutes,
        rows=board_rows,
    )


def _attendance_stats_from_board(board: AttendanceBoardRead) -> AttendanceStatsRead:
    total = len(board.rows)
    present_count = sum(1 for row in board.rows if row.status == "present")
    late_count = sum(1 for row in board.rows if row.status == "late")
    absent_count = sum(1 for row in board.rows if row.status == "absent")
    pending_count = sum(1 for row in board.rows if row.status == "pending")

    def rate(count: int) -> float:
        return round(count * 100 / total, 2) if total else 0.0

    return AttendanceStatsRead(
        session_id=board.session_id,
        class_id=board.class_id,
        total_students=total,
        present_count=present_count,
        late_count=late_count,
        absent_count=absent_count,
        pending_count=pending_count,
        present_rate=rate(present_count),
        late_rate=rate(late_count),
        absent_rate=rate(absent_count),
    )


def _attendance_request_read(db: Session, class_session: ClassSession) -> AttendanceRequestRead:
    class_ = db.get(Class, class_session.class_id)
    student_count = db.scalar(
        select(func.count(ClassEnrollment.id)).where(ClassEnrollment.class_id == class_session.class_id)
    )
    return AttendanceRequestRead(
        session_id=class_session.id,
        class_id=class_session.class_id,
        class_code=class_.code if class_ is not None else "",
        class_name=class_.name if class_ is not None else "",
        teacher_name=class_.teacher_display_name if class_ is not None else None,
        expected_start_time=class_session.expected_start_time,
        planned_end_time=class_session.planned_end_time,
        late_grace_minutes=class_session.late_grace_minutes,
        status=class_session.status,
        student_count=student_count or 0,
    )


def _session_summary(db: Session, class_session: ClassSession) -> SessionSummaryRead:
    class_ = db.get(Class, class_session.class_id)
    stats = _attendance_stats_from_board(_build_attendance_board(db, class_session))
    attended = stats.present_count + stats.late_count
    return SessionSummaryRead(
        id=class_session.id,
        class_id=class_session.class_id,
        class_code=class_.code if class_ else "",
        class_name=class_.name if class_ else "",
        teacher_name=class_.teacher_display_name if class_ else None,
        start_time=class_session.start_time,
        expected_start_time=class_session.expected_start_time,
        end_time=class_session.end_time,
        status=class_session.status,
        total_students=stats.total_students,
        present_count=stats.present_count,
        late_count=stats.late_count,
        absent_count=stats.absent_count,
        attendance_rate=round(attended * 100 / stats.total_students, 2) if stats.total_students else 0.0,
    )


@router.get("/demo", include_in_schema=False)
def demo_page() -> FileResponse:
    path = Path(__file__).resolve().parents[2] / "frontend" / "demo.html"
    return FileResponse(path, headers={"Cache-Control": "no-store"})


@router.get("/demo/{page}", include_in_schema=False)
def demo_subpage(page: str) -> FileResponse:
    pages = {"login", "students", "faces", "sessions", "attendance"}
    if page not in pages:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Page not found")
    path = Path(__file__).resolve().parents[2] / "frontend" / f"{page}.html"
    return FileResponse(path, headers={"Cache-Control": "no-store"})


@router.get("/demo/assets/{filename}", include_in_schema=False)
def demo_asset(filename: str) -> FileResponse:
    if filename not in {"app.css", "app.js"}:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Asset not found")
    path = Path(__file__).resolve().parents[2] / "frontend" / filename
    return FileResponse(path, headers={"Cache-Control": "no-store"})


@router.get("/health")
def health() -> dict[str, str]:
    return {"status": "ok"}


@router.get("/dashboard/summary", response_model=DashboardSummaryRead)
def dashboard_summary(
    _user: User = Depends(require_admin),
    db: Session = Depends(get_db),
) -> DashboardSummaryRead:
    total_students = db.scalar(select(func.count(Student.id)).where(Student.deleted_at.is_(None))) or 0
    enrolled_students = db.scalar(
        select(func.count(Student.id)).where(
            Student.deleted_at.is_(None), Student.id.in_(_completed_face_students()),
        )
    ) or 0
    active_classes = db.scalar(select(func.count(Class.id))) or 0
    open_sessions = db.scalar(
        select(func.count(ClassSession.id)).where(ClassSession.status == SessionStatus.OPEN.value)
    ) or 0
    today_start = datetime.combine(datetime.now(timezone.utc).date(), time.min, tzinfo=timezone.utc)
    attendance_today = db.scalar(
        select(func.count(AttendanceRecord.id)).where(AttendanceRecord.recorded_at >= today_start)
    ) or 0
    expected_today = db.scalar(
        select(func.count(ClassEnrollment.id))
        .join(ClassSession, ClassSession.class_id == ClassEnrollment.class_id)
        .where(ClassSession.start_time >= today_start)
    ) or 0
    return DashboardSummaryRead(
        total_students=total_students,
        enrolled_students=enrolled_students,
        active_classes=active_classes,
        open_sessions=open_sessions,
        attendance_today=attendance_today,
        attendance_rate_today=round(attendance_today * 100 / expected_today, 2) if expected_today else 0.0,
    )


@router.get("/sessions", response_model=list[SessionSummaryRead])
def list_sessions(
    limit: int = 20,
    _user: User = Depends(require_teacher_or_admin),
    db: Session = Depends(get_db),
) -> list[SessionSummaryRead]:
    safe_limit = max(1, min(limit, 100))
    sessions = db.scalars(select(ClassSession).where(ClassSession.class_id.in_(allowed_class_ids(_user)))
                          .order_by(ClassSession.start_time.desc()).limit(safe_limit)).all()
    return [_session_summary(db, item) for item in sessions]


def login(username: str, password: str, db: Session = Depends(get_db)) -> dict[str, int | str]:
    user = db.scalar(select(User).where(User.username == username))
    if user is None or not user.is_active or not verify_password(password, user.password_hash):
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Invalid credentials")
    if "$" not in user.password_hash:
        user.password_hash = hash_password(password)
    return issue_session(db, user)


@router.post("/auth/login-json")
@router.post("/auth/login")
def login_json(payload: LoginRequest, db: Session = Depends(get_db)) -> dict[str, int | str]:
    return login(payload.username, payload.password, db)


@router.post("/auth/logout", status_code=204)
def logout(_user: User = Depends(get_current_user), credentials=Depends(bearer),
           db: Session = Depends(get_db)) -> Response:
    from backend.app.models import AuthSession

    session = db.get(AuthSession, token_digest(credentials.credentials))
    if session is not None:
        db.delete(session)
    db.commit()
    return Response(status_code=204)


@router.post("/users", response_model=UserRead, dependencies=[Depends(require_admin)])
def create_user(payload: UserCreate, db: Session = Depends(get_db)) -> User:
    user = User(
        username=payload.username,
        password_hash=hash_password(payload.password),
        role=payload.role,
        full_name=payload.full_name,
    )
    db.add(user)
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Username already exists") from exc
    db.refresh(user)
    return user


@router.post("/students", response_model=StudentRead, dependencies=[Depends(require_admin)])
def create_student(payload: StudentCreate, db: Session = Depends(get_db)) -> Student:
    student = Student(**payload.model_dump())
    db.add(student)
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Student code already exists") from exc
    db.refresh(student)
    return student


@router.get("/students", response_model=list[StudentRead], dependencies=[Depends(require_admin)])
def list_students(db: Session = Depends(get_db)) -> list[Student]:
    return list(db.scalars(select(Student).where(Student.deleted_at.is_(None))).all())


@router.post("/students/import", response_model=StudentImportResult)
def import_student_file(file: UploadFile = File(...), class_id: int | None = Form(None),
                        dry_run: str = Form("true"), expected_digest: str | None = Form(None),
                        user: User = Depends(require_admin), db: Session = Depends(get_db)):
    from backend.app.services.student_import import import_students, MAX_BYTES

    if dry_run not in {"true", "false"}:
        raise HTTPException(422, "dry_run phải là true hoặc false.")
    content = file.file.read(MAX_BYTES + 1)
    return import_students(db, filename=file.filename, content=content, actor_id=user.id,
                           class_id=class_id, dry_run=dry_run == "true", expected_digest=expected_digest)


@router.get("/students/import-template", dependencies=[Depends(require_admin)])
def student_import_template(format: str = "csv") -> Response:
    headers = ["student_code", "full_name", "class_label", "consent_status"]
    example = ["00123456", "Nguyễn Văn An", "KTPM", "pending"]
    if format == "csv":
        output = StringIO(newline="")
        writer = csv.writer(output)
        writer.writerows([headers, example])
        data = output.getvalue().encode("utf-8-sig")
        media_type, extension = "text/csv; charset=utf-8", "csv"
    elif format == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sinh vien"
        sheet.append(headers)
        sheet.append(example)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="087F75")
        for column, width in (("A", 22), ("B", 32), ("C", 22), ("D", 22)):
            sheet.column_dimensions[column].width = width
        sheet["A2"].number_format = "@"
        sheet.column_dimensions["A"].number_format = "@"
        sheet.freeze_panes = "A2"
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        data = output.getvalue()
        media_type, extension = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"
    else:
        raise HTTPException(422, "Chọn định dạng csv hoặc xlsx.")
    return Response(content=data, media_type=media_type,
                    headers={"Content-Disposition": f'attachment; filename="mau-sinh-vien.{extension}"'})


@router.patch("/students/{student_id}", response_model=StudentRead)
def update_student(
    student_id: int,
    payload: StudentUpdate,
    _user: User = Depends(require_admin),
    db: Session = Depends(get_db),
) -> Student:
    student = db.get(Student, student_id)
    if student is None or student.deleted_at is not None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Student not found")

    updates = payload.model_dump(exclude_unset=True)
    for field, value in updates.items():
        setattr(student, field, value)
    if updates.get("consent_status") in {ConsentStatus.REVOKED.value, ConsentStatus.PENDING.value}:
        db.query(FaceEnrollmentDraft).filter(FaceEnrollmentDraft.student_id == student_id).delete()

    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Student code already exists") from exc
    db.refresh(student)
    return student


@router.delete("/students/{student_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_student(
    student_id: int,
    _user: User = Depends(require_admin),
    db: Session = Depends(get_db),
) -> Response:
    student = db.get(Student, student_id)
    if student is None or student.deleted_at is not None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Student not found")

    student.deleted_at = utc_now()
    db.query(FaceTemplate).filter(FaceTemplate.student_id == student_id).update({"is_active": False})
    db.query(FaceEnrollmentDraft).filter(FaceEnrollmentDraft.student_id == student_id).delete()
    db.commit()
    return Response(status_code=status.HTTP_204_NO_CONTENT)


@router.get("/admin/students/faces", response_model=list[StudentFaceSummary])
def list_students_with_faces(
    _user: User = Depends(require_admin),
    db: Session = Depends(get_db),
) -> list[StudentFaceSummary]:
    students = db.scalars(select(Student).where(Student.deleted_at.is_(None)).order_by(Student.student_code)).all()
    summaries: list[StudentFaceSummary] = []
    for student in students:
        templates = db.scalars(
            select(FaceTemplate)
            .where(FaceTemplate.student_id == student.id, FaceTemplate.is_active.is_(True))
            .order_by(FaceTemplate.created_at.desc())
        ).all()
        latest = templates[0] if templates else None
        profiles: dict[str, set[str]] = {}
        for item in templates:
            if item.pose_label in {"front", "left", "right", "up", "down"}:
                profiles.setdefault(item.model_version, set()).add(item.pose_label)
        captured_poses = max(profiles.values(), key=len, default=set())
        summaries.append(
            StudentFaceSummary(
                id=student.id,
                student_code=student.student_code,
                full_name=student.full_name,
                class_label=student.class_label,
                consent_status=student.consent_status,
                face_template_count=len(templates),
                captured_pose_count=len(captured_poses),
                enrollment_complete=len(captured_poses) == 5,
                latest_image_path=latest.image_path if latest is not None else None,
                latest_image_url=f"/admin/face-templates/{latest.id}/image" if latest is not None else None,
            )
        )
    return summaries


@router.get("/admin/face-templates/{template_id}/image")
def get_face_template_image(
    template_id: int,
    _user: User = Depends(require_admin),
    db: Session = Depends(get_db),
) -> FileResponse:
    template = db.get(FaceTemplate, template_id)
    if template is None or template.image_path is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Face image not found")
    image_path = Path(template.image_path)
    if not image_path.exists():
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Face image file missing")
    return FileResponse(image_path)


@router.get("/classes", response_model=list[ClassRead])
def list_classes(db: Session = Depends(get_db), user: User = Depends(require_teacher_or_admin)) -> list[Class]:
    return list(db.scalars(select(Class).where(Class.id.in_(allowed_class_ids(user)))).all())


@router.post("/classes", response_model=ClassRead, dependencies=[Depends(require_admin)])
def create_class(payload: ClassCreate, db: Session = Depends(get_db)) -> Class:
    class_ = Class(**payload.model_dump())
    db.add(class_)
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Class code already exists") from exc
    db.refresh(class_)
    return class_


@router.put("/classes/{class_id}/devices/{user_id}", status_code=204, dependencies=[Depends(require_admin)])
def grant_device_class_access(class_id: int, user_id: int, db: Session = Depends(get_db)) -> Response:
    from backend.app.models import DeviceClassAccess

    user = db.get(User, user_id)
    if user is None or user.role != "user" or not user.is_active or db.get(Class, class_id) is None:
        raise HTTPException(404, "Active device account or class not found")
    if db.get(DeviceClassAccess, (user_id, class_id)) is None:
        db.add(DeviceClassAccess(user_id=user_id, class_id=class_id))
        try:
            db.commit()
        except IntegrityError:
            db.rollback()
    return Response(status_code=204)


@router.delete("/classes/{class_id}/devices/{user_id}", status_code=204, dependencies=[Depends(require_admin)])
def revoke_device_class_access(class_id: int, user_id: int, db: Session = Depends(get_db)) -> Response:
    from backend.app.models import DeviceClassAccess

    access = db.get(DeviceClassAccess, (user_id, class_id))
    if access is not None:
        db.delete(access)
        db.commit()
    return Response(status_code=204)


@router.post(
    "/classes/{class_id}/enrollments/{student_id}",
    status_code=status.HTTP_204_NO_CONTENT,
    dependencies=[Depends(require_teacher_or_admin)],
)
def enroll_student(class_id: int, student_id: int, db: Session = Depends(get_db)) -> Response:
    if db.get(Class, class_id) is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Class not found")
    if db.get(Student, student_id) is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Student not found")

    db.add(ClassEnrollment(class_id=class_id, student_id=student_id))
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
    return Response(status_code=status.HTTP_204_NO_CONTENT)


@router.post("/classes/{class_id}/sessions", response_model=SessionRead)
def open_class_session(
    class_id: int,
    payload: SessionCreate,
    user: User = Depends(require_teacher_or_admin),
    db: Session = Depends(get_db),
) -> ClassSession:
    if db.get(Class, class_id) is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Class not found")
    return open_session(
        db,
        class_id=class_id,
        opened_by=user.id,
        start_time=payload.start_time,
        expected_start_time=payload.expected_start_time,
        planned_end_time=payload.planned_end_time,
        late_grace_minutes=payload.late_grace_minutes,
    )


@router.post("/attendance-requests", response_model=AttendanceRequestRead)
def create_attendance_request(
    payload: AttendanceRequestCreate,
    user: User = Depends(require_admin),
    db: Session = Depends(get_db),
) -> AttendanceRequestRead:
    class_ = db.scalar(select(Class).where(Class.code == payload.class_code))
    if class_ is None:
        class_ = Class(
            code=payload.class_code,
            name=payload.class_name,
            teacher_id=user.id,
            teacher_display_name=payload.teacher_name,
        )
        db.add(class_)
        db.flush()
    else:
        class_.name = payload.class_name
        class_.teacher_display_name = payload.teacher_name

    for item in payload.students:
        student = db.scalar(select(Student).where(Student.student_code == item.student_code))
        if student is None:
            student = Student(
                student_code=item.student_code,
                full_name=item.full_name,
                class_label=item.class_label or payload.class_code,
                consent_status=ConsentStatus.GRANTED.value,
            )
            db.add(student)
            db.flush()
        else:
            student.full_name = item.full_name
            student.class_label = item.class_label or student.class_label
            if student.consent_status == "pending":
                student.consent_status = ConsentStatus.GRANTED.value

        enrollment = db.scalar(
            select(ClassEnrollment).where(
                ClassEnrollment.class_id == class_.id,
                ClassEnrollment.student_id == student.id,
            )
        )
        if enrollment is None:
            db.add(ClassEnrollment(class_id=class_.id, student_id=student.id))

    existing_open_session = db.scalar(
        select(ClassSession).where(
            ClassSession.class_id == class_.id,
            ClassSession.status == SessionStatus.OPEN.value,
        )
    )
    if existing_open_session is not None:
        existing_open_session.expected_start_time = payload.expected_start_time
        existing_open_session.planned_end_time = payload.planned_end_time
        existing_open_session.late_grace_minutes = payload.late_grace_minutes
        db.commit()
        db.refresh(existing_open_session)
        return _attendance_request_read(db, existing_open_session)

    db.commit()
    class_session = open_session(
        db,
        class_id=class_.id,
        opened_by=user.id,
        expected_start_time=payload.expected_start_time,
        planned_end_time=payload.planned_end_time,
        late_grace_minutes=payload.late_grace_minutes,
    )
    return _attendance_request_read(db, class_session)


@router.get("/attendance-requests/open", response_model=list[AttendanceRequestRead])
def list_open_attendance_requests(
    _user: User = Depends(require_user_or_admin),
    db: Session = Depends(get_db),
) -> list[AttendanceRequestRead]:
    sessions = db.scalars(
        select(ClassSession)
        .where(ClassSession.status == SessionStatus.OPEN.value, ClassSession.class_id.in_(allowed_class_ids(_user)))
        .order_by(ClassSession.start_time.desc())
    ).all()
    return [_attendance_request_read(db, item) for item in sessions]


@router.post("/sessions/{session_id}/close", response_model=SessionRead)
def close_class_session(
    session_id: int,
    user: User = Depends(require_teacher_or_admin),
    db: Session = Depends(get_db),
) -> ClassSession:
    return close_session(db, session_id=session_id, actor_user_id=user.id)


@router.put("/sessions/{session_id}/policy", response_model=SessionRead)
def update_session_policy(
    session_id: int,
    payload: SessionPolicyUpdate,
    _user: User = Depends(require_teacher_or_admin),
    db: Session = Depends(get_db),
) -> ClassSession:
    class_session = db.get(ClassSession, session_id)
    if class_session is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Session not found")
    class_session.expected_start_time = payload.expected_start_time
    class_session.late_grace_minutes = payload.late_grace_minutes
    db.commit()
    db.refresh(class_session)
    return class_session


@router.post(
    "/sessions/{session_id}/recognition-events",
    response_model=RecognitionEventRead,
    dependencies=[Depends(verify_camera_token)],
)
def add_recognition_event(
    session_id: int,
    payload: RecognitionEventCreate,
    db: Session = Depends(get_db),
):
    event, _attendance = create_recognition_event(db, session_id=session_id, payload=payload)
    return event


@router.get("/sessions/{session_id}/attendance", response_model=list[AttendanceRead])
def list_attendance(
    session_id: int,
    _user: User = Depends(require_teacher_or_admin),
    db: Session = Depends(get_db),
) -> list[AttendanceRecord]:
    return list(db.scalars(select(AttendanceRecord).where(AttendanceRecord.session_id == session_id)).all())


@router.get("/sessions/{session_id}/attendance-board", response_model=AttendanceBoardRead)
def attendance_board(
    session_id: int,
    _user: User = Depends(require_user_or_admin),
    db: Session = Depends(get_db),
) -> AttendanceBoardRead:
    class_session = db.get(ClassSession, session_id)
    if class_session is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Session not found")
    return _build_attendance_board(db, class_session)


@router.get("/sessions/{session_id}/face-readiness")
def session_face_readiness(session_id: int, user: User = Depends(require_user_or_admin),
                           db: Session = Depends(get_db)) -> dict:
    from backend.app.security import check_class_access
    from backend.app.services.stream_recognition import _gallery

    session = db.get(ClassSession, session_id)
    if session is None:
        raise HTTPException(404, "Session not found")
    check_class_access(db, user, session.class_id)
    total = db.scalar(select(func.count(Student.id)).join(ClassEnrollment).where(
        ClassEnrollment.class_id == session.class_id, Student.deleted_at.is_(None))) or 0
    recognizer = build_default_recognizer()
    if recognizer.distance_metric != "euclidean":
        return {"ready": 0, "total": total, "engine_ready": False}
    templates, _ = _gallery(db, session.class_id, recognizer.model_version, 128)
    return {"ready": len(templates), "total": total, "engine_ready": True}


@router.get("/sessions/{session_id}/attendance-stats", response_model=AttendanceStatsRead)
def attendance_stats(
    session_id: int,
    _user: User = Depends(require_teacher_or_admin),
    db: Session = Depends(get_db),
) -> AttendanceStatsRead:
    class_session = db.get(ClassSession, session_id)
    if class_session is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Session not found")
    return _attendance_stats_from_board(_build_attendance_board(db, class_session))


@router.get("/attendance/{attendance_id}/evidence")
def get_attendance_evidence(
    attendance_id: int,
    _user: User = Depends(require_teacher_or_admin),
    db: Session = Depends(get_db),
) -> FileResponse:
    attendance = db.get(AttendanceRecord, attendance_id)
    if attendance is None or attendance.evidence_image_path is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Attendance evidence not found")
    image_path = Path(attendance.evidence_image_path)
    if not image_path.exists():
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Attendance evidence file missing")
    return FileResponse(image_path)


@router.post("/sessions/{session_id}/scan-frame", response_model=FrameScanResponse)
def scan_frame(
    session_id: int,
    payload: FrameScanRequest,
    _user: User = Depends(require_user_or_admin),
    db: Session = Depends(get_db),
) -> FrameScanResponse:
    return _scan_frame(session_id, payload, db)


@router.post("/sessions/{session_id}/scan-stream", response_model=StreamScanResponse)
def scan_stream(
    session_id: int,
    payload: StreamScanRequest,
    user: User = Depends(require_user_or_admin),
    db: Session = Depends(get_db),
) -> StreamScanResponse:
    from backend.app.services.stream_recognition import scan_stream_frame

    return scan_stream_frame(db, session_id=session_id, payload=payload, owner=("user", user.id),
                             recognizer=build_default_recognizer(), save_evidence=_save_attendance_evidence,
                             localize_reason=_localized_face_error)


@router.post("/sessions/{session_id}/camera-scan-stream", response_model=StreamScanResponse,
             dependencies=[Depends(verify_camera_token)])
def scan_camera_stream(
    session_id: int,
    payload: StreamScanRequest,
    dry_run: bool = False,
    db: Session = Depends(get_db),
) -> StreamScanResponse:
    from backend.app.services.stream_recognition import scan_stream_frame

    return scan_stream_frame(db, session_id=session_id, payload=payload, owner=("camera",), dry_run=dry_run,
                             recognizer=build_default_recognizer(), save_evidence=_save_attendance_evidence,
                             localize_reason=_localized_face_error)


@router.post(
    "/sessions/{session_id}/camera-scan-frame", response_model=FrameScanResponse,
    dependencies=[Depends(verify_camera_token)],
)
def scan_camera_frame(
    session_id: int,
    payload: FrameScanRequest,
    dry_run: bool = False,
    stream_id: str = "",
    db: Session = Depends(get_db),
) -> FrameScanResponse:
    if len(stream_id) > 80 or len(payload.camera_id) > 80:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Camera identifier is too long")
    return _scan_frame(session_id, payload, db, dry_run=dry_run,
                       confirmation_key=(session_id, payload.camera_id, stream_id))


def _scan_frame(session_id: int, payload: FrameScanRequest, db: Session, *, dry_run: bool = False,
                confirmation_key=None) -> FrameScanResponse:
    settings = get_settings()
    confirmations = _camera_confirmations(settings.attendance_min_hits, settings.attendance_window_seconds)
    class_session = db.get(ClassSession, session_id)
    if class_session is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Session not found")
    if class_session.status != SessionStatus.OPEN.value:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Session is not open")

    recognizer = build_default_recognizer()
    if recognizer.distance_metric != "euclidean":
        raise HTTPException(503, "Thiếu engine Face ID. Cài engine theo README trong môi trường backend.")
    image_base64 = payload.image_base64
    if "," in image_base64:
        image_base64 = image_base64.split(",", 1)[1]
    try:
        image_bytes = base64.b64decode(image_base64, validate=True)
    except ValueError as exc:
        confirmations.reset(confirmation_key)
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Invalid image data") from exc

    result = recognizer.extract_single_embedding(image_bytes)
    if not result.accepted:
        confirmations.reset(confirmation_key)
        return FrameScanResponse(
            decision="unknown", message=_localized_face_error(result.reason), face_box=result.face_box,
        )

    try:
        import numpy as np
        from ai_pipeline.recognition import match_embedding
    except ModuleNotFoundError as exc:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Camera recognition dependencies are not installed. Install requirements-camera.txt.",
        ) from exc

    template_rows = db.execute(
        select(FaceTemplate)
        .join(ClassEnrollment, ClassEnrollment.student_id == FaceTemplate.student_id)
        .where(
            ClassEnrollment.class_id == class_session.class_id,
            FaceTemplate.student_id.in_(select(Student.id).where(Student.consent_status == ConsentStatus.GRANTED.value)),
            FaceTemplate.is_active.is_(True),
            FaceTemplate.model_version == recognizer.model_version,
            FaceTemplate.student_id.in_(_completed_face_students(recognizer.model_version)),
        )
    ).scalars()
    templates: dict[int, list[np.ndarray]] = {}
    for template in template_rows:
        if len(template.embedding) != template.embedding_dim * 4 or template.embedding_dim != result.embedding.size:
            continue
        vector = np.frombuffer(template.embedding, dtype=np.float32)
        if np.isfinite(vector).all():
            templates.setdefault(template.student_id, []).append(vector)

    if not templates:
        confirmations.reset(confirmation_key)
        return FrameScanResponse(
            decision="unknown", message="Lớp học chưa có đủ 5 góc Face ID cùng phiên bản đang chạy. Vui lòng đăng ký lại nếu vừa đổi cấu hình xử lý ảnh.",
            face_box=result.face_box,
        )

    settings = get_settings()
    match = match_embedding(
        result.embedding,
        templates,
        threshold=(
            settings.face_distance_threshold
            if recognizer.distance_metric == "euclidean"
            else settings.attendance_threshold
        ),
        margin=(
            settings.face_distance_margin
            if recognizer.distance_metric == "euclidean"
            else settings.attendance_margin
        ),
        metric=recognizer.distance_metric,
    )
    if match.student_id is None:
        confirmations.reset(confirmation_key)
        return FrameScanResponse(
            decision=match.decision,
            message="Độ tin cậy chưa đủ, vui lòng nhìn thẳng vào camera",
            candidate_score=match.score,
            second_score=match.second_score,
            face_box=result.face_box,
        )

    if confirmation_key is not None and not confirmations.observe(confirmation_key, match.student_id):
        return FrameScanResponse(
            decision="low_confidence", message="Đang xác nhận khuôn mặt qua nhiều khung hình",
            face_box=result.face_box,
        )

    student = db.get(Student, match.student_id)
    existing_attendance = db.scalar(
        select(AttendanceRecord).where(
            AttendanceRecord.session_id == session_id,
            AttendanceRecord.student_id == match.student_id,
        )
    )
    if existing_attendance is not None:
        return FrameScanResponse(
            decision=match.decision,
            message="Đã điểm danh trong buổi này",
            already_recorded=True,
            student_id=match.student_id,
            student_code=student.student_code if student is not None else None,
            full_name=student.full_name if student is not None else None,
            class_label=student.class_label if student is not None else None,
            class_name=class_session.class_.name if class_session.class_ is not None else None,
            candidate_score=match.score,
            second_score=match.second_score,
            attendance=existing_attendance,
            face_box=result.face_box,
        )

    if dry_run:
        return FrameScanResponse(
            decision=match.decision, message="Chạy thử: không ghi điểm danh",
            student_id=match.student_id,
            student_code=student.student_code if student else None,
            full_name=student.full_name if student else None,
            class_label=student.class_label if student else None,
            class_name=class_session.class_.name if class_session.class_ else None,
            candidate_score=match.score, second_score=match.second_score, face_box=result.face_box,
        )

    event_payload = RecognitionEventCreate(
        camera_id=payload.camera_id,
        student_id=match.student_id,
        candidate_score=match.score,
        second_score=match.second_score,
        decision=match.decision,
        model_version=recognizer.model_version,
        idempotency_key=payload.idempotency_key or str(uuid.uuid4()),
    )
    _event, attendance = create_recognition_event(db, session_id=session_id, payload=event_payload)
    if attendance is None:
        return FrameScanResponse(decision="unknown", message="Chưa lưu được điểm danh. Vui lòng kiểm tra danh sách lớp hoặc liên hệ giảng viên.",
                                 face_box=result.face_box)
    try:
        _save_attendance_evidence(db, attendance=attendance, image_bytes=image_bytes, face_box=result.face_box)
    except (OSError, ValueError):
        logger.warning("Attendance saved but face evidence could not be stored")
    message = "Đúng giờ"
    if attendance is not None and attendance.status == "late":
        message = f"Trễ {attendance.late_minutes} phút"

    return FrameScanResponse(
        decision=match.decision,
        message=message,
        student_id=student.id if student is not None else match.student_id,
        student_code=student.student_code if student is not None else None,
        full_name=student.full_name if student is not None else None,
        class_label=student.class_label if student is not None else None,
        class_name=class_session.class_.name if class_session.class_ is not None else None,
        candidate_score=match.score,
        second_score=match.second_score,
        attendance=attendance,
        face_box=result.face_box,
    )


@router.put("/sessions/{session_id}/attendance/{student_id}", response_model=AttendanceRead)
def manual_attendance_update(
    session_id: int,
    student_id: int,
    payload: AttendanceManualUpdate,
    user: User = Depends(require_teacher_or_admin),
    db: Session = Depends(get_db),
) -> AttendanceRecord:
    return update_attendance_manually(
        db,
        session_id=session_id,
        student_id=student_id,
        actor_user_id=user.id,
        payload=payload,
    )


@router.post("/students/{student_id}/face-templates", response_model=FaceTemplateRead)
def register_face_template(
    student_id: int,
    file: UploadFile = File(...),
    _user: User = Depends(require_teacher_or_admin),
    db: Session = Depends(get_db),
) -> object:
    recognizer = build_default_recognizer()
    return register_face_image(db, student_id=student_id, file=file, recognizer=recognizer)


def _enrollment_progress(db: Session, student_id: int, model_version: str | None = None) -> tuple[list[str], list[str]]:
    from ai_pipeline.recognition import ENROLLMENT_POSES

    captured = set(
        db.scalars(
            select(FaceTemplate.pose_label).where(
                FaceTemplate.student_id == student_id,
                FaceTemplate.is_active.is_(True),
                FaceTemplate.pose_label.is_not(None),
                *([FaceTemplate.model_version == model_version] if model_version is not None else []),
            )
        ).all()
    )
    ordered = [pose for pose in ENROLLMENT_POSES if pose in captured]
    return ordered, [pose for pose in ENROLLMENT_POSES if pose not in captured]


def _completed_face_profiles():
    return (
        select(FaceTemplate.student_id, FaceTemplate.model_version)
        .where(FaceTemplate.is_active.is_(True), FaceTemplate.pose_label.in_(("front", "left", "right", "up", "down")))
        .group_by(FaceTemplate.student_id, FaceTemplate.model_version)
        .having(func.count(func.distinct(FaceTemplate.pose_label)) >= 5)
    )


def _completed_face_students(model_version: str | None = None):
    profiles = _completed_face_profiles().subquery()
    query = select(profiles.c.student_id).distinct()
    return query.where(profiles.c.model_version == model_version) if model_version is not None else query


def _pipeline_status(*, load_engine: bool = True) -> dict:
    settings = get_settings()
    recognizer = build_default_recognizer() if load_engine else None
    analysis = getattr(recognizer, "_analysis", None)
    face_detector = getattr(recognizer, "_face_detector", None)
    options = analysis.preprocessing if analysis else None
    candidate = analysis.shadow_options() if analysis else None

    return {
        "engine": recognizer.model_name if recognizer else None,
        "model_version": recognizer.model_version if recognizer else None,
        "face_detection": face_detector.status() if face_detector else {"backend": "hog", "yunet_ready": False},
        "preprocessing_mode": settings.preprocessing_mode,
        "stream_recognition": True,
        "stream_max_faces_per_frame": settings.stream_max_faces_per_frame,
        "stream_confirmation_hits": settings.attendance_min_hits,
        "stream_min_observation_interval_seconds": settings.stream_min_observation_interval_seconds,
        "stream_min_confirmation_seconds": settings.stream_min_confirmation_seconds,
        "attendance_engine_ready": bool(recognizer and recognizer.distance_metric == "euclidean"),
        "preprocessing_version": settings.preprocessing_version,
        "preprocessing_profile": options.profile if options else None,
        "shadow_candidate_profile": candidate.profile if candidate else None,
        "shadow_sample_rate": options.shadow_sample_rate if candidate else 0,
        "metrics": analysis.metrics() if analysis else None,
    }


@router.get("/face-analysis/status")
def face_analysis_status(_user: User = Depends(require_user_or_admin)) -> dict:
    return _pipeline_status()


@router.get("/camera/pipeline", dependencies=[Depends(verify_camera_token)])
def camera_pipeline_status() -> dict:
    settings = get_settings()
    # Legacy local workers can use a backend without any camera dependencies.
    return _pipeline_status(load_engine=settings.preprocessing_mode != "off")


@router.get("/face-enrollment/engine")
def face_enrollment_engine(_user: User = Depends(require_teacher_or_admin)) -> dict[str, str | bool]:
    recognizer = build_default_recognizer()
    return {"available": recognizer.supports_pose, "engine": recognizer.model_name}


@router.post("/students/{student_id}/face-enrollment/start")
def start_face_enrollment(student_id: int, user: User = Depends(require_teacher_or_admin),
                          db: Session = Depends(get_db)) -> dict[str, str]:
    student = db.get(Student, student_id)
    if student is None or student.deleted_at is not None:
        raise HTTPException(404, "Không tìm thấy sinh viên")
    if student.consent_status != ConsentStatus.GRANTED.value:
        raise HTTPException(409, "Sinh viên chưa đồng ý sử dụng Face ID")
    db.query(FaceEnrollmentDraft).filter(FaceEnrollmentDraft.student_id == student_id).delete()
    enrollment_id = uuid.uuid4().hex
    # A marker makes an empty session valid and invalidates any earlier session.
    db.add(FaceEnrollmentDraft(student_id=student_id, enrollment_id=enrollment_id, owner_id=user.id,
                               pose_label="session", model_name="", model_version="", embedding=b"",
                               embedding_dim=0, image_bytes=b""))
    db.commit()
    return {"enrollment_id": enrollment_id}


@router.delete("/students/{student_id}/face-enrollment/draft", status_code=status.HTTP_204_NO_CONTENT)
def cancel_face_enrollment(student_id: int, enrollment_id: str,
                           user: User = Depends(require_teacher_or_admin), db: Session = Depends(get_db)) -> Response:
    db.query(FaceEnrollmentDraft).filter(FaceEnrollmentDraft.student_id == student_id,
        FaceEnrollmentDraft.enrollment_id == enrollment_id, FaceEnrollmentDraft.owner_id == user.id).delete()
    db.commit()
    return Response(status_code=status.HTTP_204_NO_CONTENT)


@router.post("/students/{student_id}/face-enrollment/reset", status_code=status.HTTP_204_NO_CONTENT)
def reset_face_enrollment(
    student_id: int,
    user: User = Depends(require_teacher_or_admin),
    db: Session = Depends(get_db),
) -> Response:
    student = db.get(Student, student_id)
    if student is None or student.deleted_at is not None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Không tìm thấy sinh viên")
    changed = db.query(FaceTemplate).filter(
        FaceTemplate.student_id == student_id,
        FaceTemplate.is_active.is_(True),
    ).update({"is_active": False})
    db.query(FaceEnrollmentDraft).filter(FaceEnrollmentDraft.student_id == student_id).delete()
    write_audit(
        db,
        actor_user_id=user.id,
        action="face_enrollment.reset",
        entity_type="Student",
        entity_id=student_id,
        after={"templates_deactivated": changed},
    )
    db.commit()
    return Response(status_code=status.HTTP_204_NO_CONTENT)


@router.post("/students/{student_id}/face-enrollment/scan", response_model=FaceEnrollmentProgress)
def scan_face_enrollment(
    student_id: int,
    payload: FaceEnrollmentScanRequest,
    _user: User = Depends(require_teacher_or_admin),
    db: Session = Depends(get_db),
) -> FaceEnrollmentProgress:
    from ai_pipeline.recognition import ENROLLMENT_POSES, classify_enrollment_pose

    if payload.target_pose not in ENROLLMENT_POSES:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Góc quét không hợp lệ")
    student = db.get(Student, student_id)
    if student is None or student.deleted_at is not None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Không tìm thấy sinh viên")
    if student.consent_status != ConsentStatus.GRANTED.value:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Sinh viên chưa đồng ý sử dụng Face ID")

    if payload.enrollment_id:
        marker = db.scalar(select(FaceEnrollmentDraft).where(
            FaceEnrollmentDraft.student_id == student_id, FaceEnrollmentDraft.pose_label == "session",
            FaceEnrollmentDraft.enrollment_id == payload.enrollment_id, FaceEnrollmentDraft.owner_id == _user.id))
        if marker is None:
            raise HTTPException(409, "Lượt đăng ký đã hết hiệu lực. Vui lòng bắt đầu lại.")

    recognizer = build_default_recognizer()
    if not recognizer.supports_pose:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Chưa cài face_recognition. Xem hướng dẫn cài engine Face ID trong README.",
        )
    raw = payload.image_base64.split(",", 1)[-1]
    try:
        image_bytes = base64.b64decode(raw, validate=True)
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Dữ liệu ảnh không hợp lệ") from exc

    result = recognizer.extract_single_embedding(
        image_bytes,
        num_jitters=get_settings().enrollment_num_jitters,
    )
    detected = (
        classify_enrollment_pose(result.pose, payload.target_pose, payload.neutral_pitch)
        if result.pose else None
    )
    if payload.enrollment_id:
        drafts = db.scalars(select(FaceEnrollmentDraft).where(
            FaceEnrollmentDraft.student_id == student_id, FaceEnrollmentDraft.enrollment_id == payload.enrollment_id,
            FaceEnrollmentDraft.pose_label.in_(ENROLLMENT_POSES))).all()
        if any(item.model_version != recognizer.model_version for item in drafts):
            raise HTTPException(409, "Cấu hình Face ID đã thay đổi. Vui lòng bắt đầu lại.")
        captured = [pose for pose in ENROLLMENT_POSES if any(item.pose_label == pose for item in drafts)]
        missing = [pose for pose in ENROLLMENT_POSES if pose not in captured]
    else:
        captured, missing = _enrollment_progress(db, student_id, recognizer.model_version)
    base_response = {
        "detected_pose": detected,
        "target_pose": payload.target_pose,
        "yaw": result.pose.yaw if result.pose else None,
        "pitch": result.pose.pitch if result.pose else None,
        "captured_poses": captured,
        "missing_poses": missing,
        "progress_percent": len(captured) * 20,
        "completed": not missing,
        "engine": recognizer.model_name,
    }
    if not result.accepted:
        return FaceEnrollmentProgress(accepted=False, message=_localized_face_error(result.reason), **base_response)
    if detected is None:
        return FaceEnrollmentProgress(
            accepted=False, message="Chưa xác định được hướng khuôn mặt. Vui lòng nhìn rõ vào camera và thử lại.",
            **base_response,
        )
    if detected != payload.target_pose:
        names = {"front": "chính diện", "left": "trái", "right": "phải", "up": "lên trên", "down": "xuống dưới"}
        return FaceEnrollmentProgress(
            accepted=False,
            message=f"Đang thấy góc {names.get(detected, detected)}; hãy giữ đúng hướng yêu cầu",
            **base_response,
        )

    if payload.enrollment_id:
        if payload.target_pose not in captured:
            db.add(FaceEnrollmentDraft(student_id=student_id, enrollment_id=payload.enrollment_id,
                owner_id=_user.id, pose_label=payload.target_pose, model_name=recognizer.model_name,
                model_version=recognizer.model_version, embedding=result.embedding.astype("float32").tobytes(),
                embedding_dim=int(result.embedding.shape[0]), image_bytes=image_bytes, quality_score=result.quality_score))
            db.flush()
        drafts = db.scalars(select(FaceEnrollmentDraft).where(
            FaceEnrollmentDraft.student_id == student_id, FaceEnrollmentDraft.enrollment_id == payload.enrollment_id,
            FaceEnrollmentDraft.pose_label.in_(ENROLLMENT_POSES))).all()
        captured = [pose for pose in ENROLLMENT_POSES if any(item.pose_label == pose for item in drafts)]
        missing = [pose for pose in ENROLLMENT_POSES if pose not in captured]
        if not missing:
            # Switch the entire set in one transaction; attendance never sees a partial set.
            from ai_pipeline.recognition import EmbeddingResult
            import numpy as np

            db.query(FaceTemplate).filter(FaceTemplate.student_id == student_id,
                                          FaceTemplate.is_active.is_(True)).update({"is_active": False})
            for draft in drafts:
                vector = np.frombuffer(draft.embedding, dtype=np.float32)
                if vector.shape != (draft.embedding_dim,) or not np.isfinite(vector).all():
                    raise HTTPException(409, "Mẫu khuôn mặt không hợp lệ. Vui lòng bắt đầu lại.")
                register_face_bytes(db, student_id=student_id, image_bytes=draft.image_bytes,
                    filename=f"faceid-{draft.pose_label}-{uuid.uuid4()}.jpg", recognizer=recognizer,
                    pose_label=draft.pose_label,
                    extracted_result=EmbeddingResult(True, vector, draft.quality_score), commit=False)
            db.query(FaceEnrollmentDraft).filter(FaceEnrollmentDraft.student_id == student_id,
                                                  FaceEnrollmentDraft.enrollment_id == payload.enrollment_id).delete()
        db.commit()
        return FaceEnrollmentProgress(accepted=True, message="Đã ghi nhận góc khuôn mặt", detected_pose=detected,
            target_pose=payload.target_pose, yaw=result.pose.yaw if result.pose else None,
            pitch=result.pose.pitch if result.pose else None, captured_poses=captured, missing_poses=missing,
            progress_percent=len(captured) * 20, completed=not missing, engine=recognizer.model_name)

    existing = db.scalar(
        select(FaceTemplate).where(
            FaceTemplate.student_id == student_id,
            FaceTemplate.pose_label == payload.target_pose,
            FaceTemplate.model_version == recognizer.model_version,
            FaceTemplate.is_active.is_(True),
        )
    )
    if existing is None:
        register_face_bytes(
            db,
            student_id=student_id,
            image_bytes=image_bytes,
            filename=f"faceid-{payload.target_pose}-{uuid.uuid4()}.jpg",
            recognizer=recognizer,
            pose_label=payload.target_pose,
            extracted_result=result,
        )
    captured, missing = _enrollment_progress(db, student_id, recognizer.model_version)
    return FaceEnrollmentProgress(
        accepted=True,
        message="Đã ghi nhận góc khuôn mặt",
        detected_pose=detected,
        target_pose=payload.target_pose,
        yaw=result.pose.yaw if result.pose else None,
        pitch=result.pose.pitch if result.pose else None,
        captured_poses=captured,
        missing_poses=missing,
        progress_percent=len(captured) * 20,
        completed=not missing,
        engine=recognizer.model_name,
    )


def _student_code_from_filename(filename: str) -> str | None:
    stem = Path(filename).stem.strip()
    if not stem:
        return None
    for separator in ("_", "-", " "):
        if separator in stem:
            stem = stem.split(separator, 1)[0]
            break
    return stem.strip() or None


@router.post("/admin/face-import", response_model=FaceImportResponse)
def import_student_faces(
    files: list[UploadFile] = File(...),
    class_id: int | None = None,
    _user: User = Depends(require_admin),
    db: Session = Depends(get_db),
) -> FaceImportResponse:
    recognizer = build_default_recognizer()
    items: list[FaceImportItem] = []
    imported = 0

    for file in files:
        filename = file.filename or "unknown.jpg"
        student_code = _student_code_from_filename(filename)
        if student_code is None:
            items.append(FaceImportItem(filename=filename, status="failed", detail="Cannot parse student code"))
            continue

        student = db.scalar(select(Student).where(Student.student_code == student_code))
        if student is None:
            student = Student(
                student_code=student_code,
                full_name=student_code,
                class_label=None,
                consent_status=ConsentStatus.GRANTED.value,
            )
            db.add(student)
            db.flush()

        if class_id is not None:
            if db.get(Class, class_id) is None:
                db.rollback()
                raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Class not found")
            enrollment = db.scalar(
                select(ClassEnrollment).where(
                    ClassEnrollment.class_id == class_id,
                    ClassEnrollment.student_id == student.id,
                )
            )
            if enrollment is None:
                db.add(ClassEnrollment(class_id=class_id, student_id=student.id))

        image_bytes = file.file.read()
        try:
            register_face_bytes(
                db,
                student_id=student.id,
                image_bytes=image_bytes,
                filename=f"{uuid.uuid4()}-{filename}",
                recognizer=recognizer,
            )
        except HTTPException as exc:
            db.rollback()
            items.append(
                FaceImportItem(
                    filename=filename,
                    student_code=student_code,
                    status="failed",
                    detail=str(exc.detail),
                )
            )
            continue

        imported += 1
        items.append(
            FaceImportItem(
                filename=filename,
                student_code=student_code,
                status="imported",
                detail="OK",
            )
        )

    return FaceImportResponse(imported=imported, failed=len(items) - imported, items=items)


@router.get("/classes/{class_id}/face-templates", dependencies=[Depends(verify_camera_token)])
def list_class_face_templates(class_id: int, db: Session = Depends(get_db)) -> list[dict[str, object]]:
    rows = db.execute(
        select(FaceTemplate)
        .join(ClassEnrollment, ClassEnrollment.student_id == FaceTemplate.student_id)
        .where(
            ClassEnrollment.class_id == class_id,
            FaceTemplate.student_id.in_(select(Student.id).where(Student.consent_status == ConsentStatus.GRANTED.value)),
            FaceTemplate.is_active.is_(True),
            tuple_(FaceTemplate.student_id, FaceTemplate.model_version).in_(_completed_face_profiles()),
        )
    ).scalars()
    return [
        {
            "student_id": template.student_id,
            "model_name": template.model_name,
            "model_version": template.model_version,
            "embedding_dim": template.embedding_dim,
            "embedding_base64": base64.b64encode(template.embedding).decode("ascii"),
        }
        for template in rows
    ]


@router.get("/reports/attendance.csv")
def export_attendance_csv(
    session_id: int,
    _user: User = Depends(require_teacher_or_admin),
    db: Session = Depends(get_db),
) -> Response:
    rows = db.execute(
        select(
            Student.student_code,
            Student.full_name,
            AttendanceRecord.status,
            AttendanceRecord.source,
            AttendanceRecord.recorded_at,
        )
        .join(AttendanceRecord, AttendanceRecord.student_id == Student.id)
        .where(AttendanceRecord.session_id == session_id)
        .order_by(Student.student_code)
    ).all()
    output = StringIO()
    writer = csv.writer(output)
    writer.writerow(["student_code", "full_name", "status", "source", "recorded_at"])
    writer.writerows(rows)
    return Response(
        content=output.getvalue(),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="attendance-session-{session_id}.csv"'},
    )


@router.get("/reports/attendance.xlsx")
def export_attendance_excel(
    session_id: int,
    _user: User = Depends(require_teacher_or_admin),
    db: Session = Depends(get_db),
) -> Response:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ModuleNotFoundError as exc:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Excel export dependency is not installed. Install requirements.txt.",
        ) from exc

    class_session = db.get(ClassSession, session_id)
    if class_session is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Session not found")

    class_ = db.get(Class, class_session.class_id)
    board = _build_attendance_board(db, class_session)
    stats = _attendance_stats_from_board(board)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Attendance"

    title = f"Attendance report - {class_.code if class_ else ''}"
    sheet["A1"] = title
    sheet["A1"].font = Font(size=16, bold=True)
    sheet.merge_cells("A1:H1")

    metadata = [
        ("Class", f"{class_.code} - {class_.name}" if class_ else str(class_session.class_id)),
        ("Teacher", class_.teacher_display_name if class_ else ""),
        ("Expected start", class_session.expected_start_time.isoformat() if class_session.expected_start_time else ""),
        ("Late grace minutes", class_session.late_grace_minutes),
        ("Total students", stats.total_students),
        ("Present rate", f"{stats.present_rate}%"),
        ("Late rate", f"{stats.late_rate}%"),
        ("Absent rate", f"{stats.absent_rate}%"),
    ]
    for index, (label, value) in enumerate(metadata, start=3):
        sheet.cell(row=index, column=1, value=label).font = Font(bold=True)
        sheet.cell(row=index, column=2, value=value)

    summary_start = 3
    summary = [
        ("Present", stats.present_count),
        ("Late", stats.late_count),
        ("Absent", stats.absent_count),
        ("Pending", stats.pending_count),
    ]
    for offset, (label, value) in enumerate(summary):
        row = summary_start + offset
        sheet.cell(row=row, column=4, value=label).font = Font(bold=True)
        sheet.cell(row=row, column=5, value=value)

    header_row = 13
    headers = [
        "Student code",
        "Full name",
        "Class",
        "Recorded at",
        "Status",
        "Source",
        "Late minutes",
        "Evidence image",
    ]
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=column, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
        cell.alignment = Alignment(horizontal="center")

    for row_index, row in enumerate(board.rows, start=header_row + 1):
        values = [
            row.student_code,
            row.full_name,
            row.class_label or "",
            row.recorded_at.isoformat() if row.recorded_at else "",
            row.status,
            row.source or "",
            row.late_minutes,
            row.evidence_image_url or "",
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row=row_index, column=column, value=value)

    for column in range(1, len(headers) + 1):
        column_letter = get_column_letter(column)
        max_length = max(
            len(str(sheet.cell(row=row, column=column).value or ""))
            for row in range(1, sheet.max_row + 1)
        )
        sheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 42)
    sheet.freeze_panes = "A14"

    output = BytesIO()
    workbook.save(output)
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="attendance-session-{session_id}.xlsx"'},
    )
