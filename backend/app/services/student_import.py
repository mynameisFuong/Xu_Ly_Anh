"""Bounded CSV/XLSX student import with preview and no updates to existing records."""

import csv
from hashlib import sha256
from io import BytesIO, StringIO
from pathlib import Path
import re
import unicodedata
from zipfile import BadZipFile, ZipFile
from openpyxl.utils.exceptions import InvalidFileException

from fastapi import HTTPException
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError

from backend.app.models import Class, ClassEnrollment, Student
from backend.app.services.audit import write_audit

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 2000
MAX_COLUMNS = 100


def normalized(value):
    text = unicodedata.normalize("NFKD", str(value).strip().lower().replace("đ", "d"))
    return re.sub(r"[^a-z0-9]", "", "".join(char for char in text if not unicodedata.combining(char)))


ALIASES = {
    "student_code": {"studentcode", "mssv", "masinhvien", "masosinhvien"},
    "full_name": {"fullname", "hovaten", "hoten", "tensinhvien"},
    "class_label": {"classlabel", "lop", "lophanhchinh"},
    "consent_status": {"consentstatus", "trangthaidongy", "dongy"},
}
CONSENT = {"pending": "pending", "chodongy": "pending", "granted": "granted",
           "dadongy": "granted", "dongy": "granted", "revoked": "revoked", "dathuhoi": "revoked"}


def workbook_rows(content):
    try:
        with ZipFile(BytesIO(content)) as archive:
            if len(archive.infolist()) > 1000 or sum(item.file_size for item in archive.infolist()) > 25 * 1024 * 1024:
                raise HTTPException(422, "File Excel giải nén quá lớn. Hãy tách danh sách thành file nhỏ hơn.")
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            sheet = workbook.worksheets[0]
            # Do not trust inflated or incorrectly saved worksheet dimensions.
            sheet.reset_dimensions()
            for number, cells in enumerate(sheet.iter_rows(), 1):
                if number > MAX_ROWS + 1:
                    raise HTTPException(422, f"File Excel tối đa {MAX_ROWS} dòng sinh viên sau tiêu đề.")
                values = []
                for cell in cells:
                    value = cell.value
                    if cell.data_type == "f":
                        value = _InvalidCell("Không dùng công thức Excel; hãy dán giá trị vào ô.")
                    elif isinstance(value, (int, float)) and not isinstance(value, bool):
                        if int(value) == value and abs(value) < 10 ** 15:
                            pattern = str(cell.number_format)
                            value = str(int(value)).zfill(len(pattern)) if re.fullmatch(r"0{2,40}", pattern) else str(int(value))
                        else:
                            value = _InvalidCell("MSSV và dữ liệu chữ cần lưu dạng Text; không dùng số thập phân hoặc quá dài.")
                    elif value is not None and not isinstance(value, str):
                        value = _InvalidCell("Ô cần chứa văn bản, không dùng ngày tháng hoặc kiểu dữ liệu khác.")
                    values.append(value)
                yield number, values
        finally:
            workbook.close()
    except HTTPException:
        raise
    except (BadZipFile, InvalidFileException, ValueError, KeyError, IndexError, OSError,
            OverflowError, SyntaxError, RuntimeError, EOFError) as exc:
        raise HTTPException(422, "Không đọc được file Excel. Hãy lưu lại dưới dạng .xlsx.") from exc


class _InvalidCell:
    def __init__(self, message):
        self.message = message


def csv_rows(content):
    try:
        encoding = "utf-16" if content.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
        text = content.decode(encoding)
    except UnicodeError as exc:
        raise HTTPException(422, "Không đọc được tiếng Việt. Hãy lưu file CSV UTF-8 hoặc dùng Excel .xlsx.") from exc
    if "\x00" in text:
        raise HTTPException(422, "File CSV chứa dữ liệu không hợp lệ.")
    try:
        try:
            delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t").delimiter
        except csv.Error:
            delimiter = ";" if ";" in text.splitlines()[0] else ","
        reader = csv.reader(StringIO(text, newline=""), delimiter=delimiter, strict=True)
        for values in reader:
            yield reader.line_num, values
    except (csv.Error, IndexError) as exc:
        raise HTTPException(422, "CSV không đúng định dạng. Kiểm tra dấu phân cách và dấu ngoặc kép.") from exc


def parse_students(filename, content):
    if not content:
        raise HTTPException(422, "File trống. Hãy chọn file có danh sách sinh viên.")
    if len(content) > MAX_BYTES:
        raise HTTPException(413, "File tối đa 5 MB.")
    suffix = Path(filename or "").suffix.lower()
    if suffix not in {".csv", ".xlsx"}:
        raise HTTPException(422, "Chỉ hỗ trợ .csv hoặc .xlsx. Với Excel .xls, hãy lưu lại thành .xlsx.")
    source = csv_rows(content) if suffix == ".csv" else workbook_rows(content)
    mapping, results, header_width = None, [], 0
    try:
        for number, values in source:
            if not any(value is not None and str(value).strip() for value in values):
                continue
            while values and values[-1] in (None, ""):
                values.pop()
            if len(values) > MAX_COLUMNS:
                raise HTTPException(422, f"File tối đa {MAX_COLUMNS} cột.")
            if mapping is None:
                mapping = {}
                header_width = len(values)
                for index, value in enumerate(values):
                    for field, aliases in ALIASES.items():
                        if normalized(value) in aliases:
                            if field in mapping:
                                raise HTTPException(422, f"Cột {field} bị lặp trong tiêu đề.")
                            mapping[field] = index
                if not {"student_code", "full_name"} <= mapping.keys():
                    raise HTTPException(422, "Cần cột MSSV và Họ và tên (hoặc student_code, full_name) ở dòng tiêu đề.")
                continue
            if len(results) >= MAX_ROWS:
                raise HTTPException(422, f"Mỗi file tối đa {MAX_ROWS} sinh viên. Hãy tách file trước khi nhập.")
            item = dict(row_number=number, student_code="", full_name="", class_label=None,
                        consent_status="pending", status="valid", message="Sẵn sàng thêm mới")
            errors = []
            if len(values) > header_width:
                errors.append("Số cột nhiều hơn tiêu đề; kiểm tra dấu phân cách và dấu ngoặc kép.")
            for field, maximum in (("student_code", 40), ("full_name", 160), ("class_label", 80), ("consent_status", 40)):
                index = mapping.get(field)
                value = values[index] if index is not None and index < len(values) else None
                if isinstance(value, _InvalidCell):
                    errors.append(value.message)
                    continue
                value = str(value).strip() if value is not None else ""
                if any(ord(char) < 32 for char in value):
                    errors.append(f"Cột {field} chứa ký tự điều khiển hoặc xuống dòng.")
                elif value.startswith(("=", "+", "-", "@")):
                    errors.append(f"Cột {field} phải là văn bản, không bắt đầu bằng dấu công thức =, +, - hoặc @.")
                elif len(value) > maximum:
                    errors.append(f"Cột {field} tối đa {maximum} ký tự.")
                elif field == "consent_status":
                    consent = CONSENT.get(normalized(value or "pending"))
                    if consent is None:
                        errors.append("Đồng ý chỉ nhận pending, granted, revoked hoặc tên trạng thái tiếng Việt.")
                    else:
                        item[field] = consent
                else:
                    item[field] = value or (None if field == "class_label" else "")
            if not item["student_code"] or not item["full_name"]:
                errors.append("Thiếu MSSV hoặc họ và tên.")
            if errors:
                item.update(status="error", message=" ".join(dict.fromkeys(errors)))
            results.append(item)
    finally:
        source.close()
    if not results:
        raise HTTPException(422, "File không có dòng sinh viên nào sau tiêu đề.")
    return results


def import_students(db, *, filename, content, actor_id, class_id=None, dry_run=True, expected_digest=None):
    if class_id is not None and db.get(Class, class_id) is None:
        raise HTTPException(404, "Lớp học đã chọn không tồn tại.")
    digest = sha256(str(class_id).encode() + b"\0" + content).hexdigest()
    if not dry_run and expected_digest != digest:
        raise HTTPException(409, "File hoặc lớp học đã thay đổi. Hãy xem trước lại trước khi nhập.")
    rows = parse_students(filename, content)
    # Query in bounded batches to stay below SQLite's bind parameter limits.
    codes = list({item["student_code"] for item in rows if item["student_code"]})
    existing = set()
    for start in range(0, len(codes), 400):
        existing.update(db.scalars(select(Student.student_code).where(Student.student_code.in_(codes[start:start + 400]))))
    seen = set()
    for item in rows:
        if item["status"] == "error":
            continue
        code = item["student_code"]
        if code in existing or code in seen:
            item.update(status="duplicate", message="MSSV đã tồn tại hoặc bị lặp trong file; giữ nguyên hồ sơ cũ.")
        seen.add(code)
    if not dry_run:
        try:
            for item in rows:
                if item["status"] != "valid":
                    continue
                student = Student(**{field: item[field] for field in ALIASES})
                db.add(student)
                db.flush()
                if class_id is not None:
                    db.add(ClassEnrollment(class_id=class_id, student_id=student.id))
                item.update(status="imported", message="Đã thêm sinh viên" + (" và ghi danh vào lớp đã chọn." if class_id else "."))
            write_audit(db, actor_user_id=actor_id, action="students.import", entity_type="Student", entity_id="batch",
                        after={"imported": sum(row["status"] == "imported" for row in rows), "class_id": class_id})
            db.commit()
        except IntegrityError as exc:
            db.rollback()
            raise HTTPException(409, "Danh sách vừa thay đổi hoặc MSSV bị trùng. Hãy xem trước và nhập lại.") from exc
        except Exception:
            db.rollback()
            raise
    return dict(dry_run=dry_run, digest=digest, total=len(rows),
                valid=sum(item["status"] in {"valid", "imported"} for item in rows),
                imported=sum(item["status"] == "imported" for item in rows),
                skipped=sum(item["status"] == "duplicate" for item in rows),
                failed=sum(item["status"] == "error" for item in rows), rows=rows)
