from io import BytesIO

import pytest
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, event, select, func
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from backend.app.api import router
from backend.app.database import Base, get_db
from backend.app.models import AuditLog, Class, ClassEnrollment, Student, User, utc_now
from backend.app.security import issue_session
from backend.app.services.student_import import parse_students


@pytest.fixture
def import_context():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    with sessionmaker(bind=engine, expire_on_commit=False)() as db:
        admin = User(username="admin", password_hash="unused", role="admin", full_name="Admin")
        teacher = User(username="teacher", password_hash="unused", role="teacher", full_name="Teacher")
        db.add_all([admin, teacher])
        db.flush()
        class_ = Class(code="IMG101", name="Image", teacher_id=teacher.id)
        db.add(class_)
        db.commit()
        app = FastAPI()
        app.include_router(router)
        app.dependency_overrides[get_db] = lambda: db
        with TestClient(app) as client:
            client.headers["Authorization"] = f"Bearer {issue_session(db, admin)['access_token']}"
            yield client, db, class_, teacher
    engine.dispose()


def upload(client, content, *, filename="students.csv", **data):
    return client.post("/students/import", files={"file": (filename, content)}, data=data)


def test_csv_preview_then_import_valid_rows_preserves_existing_and_can_enroll(import_context):
    client, db, class_, _ = import_context
    existing = Student(student_code="OLD", full_name="Original", consent_status="revoked", deleted_at=utc_now())
    db.add(existing)
    db.commit()
    content = ("MSSV,Họ và tên,Lớp hành chính\n"
               "00123,Nguyễn Văn An,KTPM\n"
               "00123,Duplicate,KTPM\n"
               "OLD,Replacement,KTPM\n"
               "BAD,,KTPM\n").encode("utf-8-sig")
    preview = upload(client, content, class_id=class_.id)
    assert preview.status_code == 200
    result = preview.json()
    assert (result["valid"], result["skipped"], result["failed"], result["imported"]) == (1, 2, 1, 0)
    assert db.scalar(select(func.count(Student.id))) == 1
    assert db.scalar(select(AuditLog)) is None
    response = upload(client, content, class_id=class_.id, dry_run="false", expected_digest=result["digest"])
    assert response.status_code == 200 and response.json()["imported"] == 1
    student = db.scalar(select(Student).where(Student.student_code == "00123"))
    assert student.full_name == "Nguyễn Văn An" and student.consent_status == "pending"
    assert db.scalar(select(ClassEnrollment)).student_id == student.id
    assert existing.full_name == "Original" and existing.consent_status == "revoked" and existing.deleted_at is not None
    assert db.scalar(select(AuditLog)).action == "students.import"
    assert upload(client, content).json()["valid"] == 0


def test_preview_binding_and_role_checks_prevent_unreviewed_or_unauthorized_import(import_context):
    client, db, class_, teacher = import_context
    content = b"student_code,full_name\nA1,Person One\n"
    digest = upload(client, content).json()["digest"]
    assert upload(client, content, dry_run="false").status_code == 409
    assert upload(client, content, dry_run="false", expected_digest=digest, class_id=class_.id).status_code == 409
    assert upload(client, content + b"A2,Person Two\n", dry_run="false", expected_digest=digest).status_code == 409
    assert db.scalar(select(Student)) is None
    client.headers["Authorization"] = f"Bearer {issue_session(db, teacher)['access_token']}"
    assert upload(client, content).status_code == 403
    assert client.get("/students/import-template").status_code == 403
    client.headers.clear()
    assert upload(client, content).status_code == 401


def test_csv_supports_utf16_semicolon_and_rejects_invalid_headers_and_status():
    data = 'Mã số sinh viên;Họ và tên;Trạng thái đồng ý\n0007;"Nguyễn, Văn An";Đã đồng ý\n'
    rows = parse_students("data.CSV", data.encode("utf-16"))
    assert rows[0]["student_code"] == "0007" and rows[0]["full_name"] == "Nguyễn, Văn An"
    assert rows[0]["consent_status"] == "granted"
    rows = parse_students("data.csv", b"student_code,full_name,consent_status\nA,Name,unknown\n")
    assert rows[0]["status"] == "error"
    formula = parse_students("data.csv", b"student_code,full_name\nA,=1+2\n")
    assert formula[0]["status"] == "error"
    for content in (b"wrong,columns\nA,Name\n", b"student_code,mssv,full_name\nA,A,Name\n", b"", b"\x80\x81"):
        with pytest.raises(HTTPException):
            parse_students("data.csv", content)


def test_excel_preserves_text_and_zero_formats_and_rejects_formula_rows():
    book = Workbook()
    sheet = book.active
    sheet.append(["student_code", "full_name", "class_label"])
    sheet.append([123, "Numeric code", "KTPM"])
    sheet["A2"].number_format = "000000"
    sheet.append(["00001234", "Text code", "KTPM"])
    sheet.append(["=1+2", "Formula", "KTPM"])
    sheet.append([1234567890123456, "Too long numeric", "KTPM"])
    output = BytesIO()
    book.save(output)
    book.close()
    rows = parse_students("data.xlsx", output.getvalue())
    assert rows[0]["student_code"] == "000123" and rows[1]["student_code"] == "00001234"
    assert all(row["status"] == "error" for row in rows[2:])


def test_templates_round_trip_with_valid_headers_unicode_and_leading_zeros(import_context):
    client, _, _, _ = import_context
    for format in ("csv", "xlsx"):
        response = client.get(f"/students/import-template?format={format}")
        assert response.status_code == 200 and "attachment" in response.headers["content-disposition"]
        row = parse_students(f"template.{format}", response.content)[0]
        assert row["student_code"] == "00123456" and row["consent_status"] == "pending"
        assert row["full_name"] == "Nguyễn Văn An" and row["status"] == "valid"
        if format == "xlsx":
            book = load_workbook(BytesIO(response.content))
            assert book.active["A2"].number_format == "@"
            book.close()


def test_database_failure_rolls_back_all_new_rows_and_enrollments(import_context):
    client, db, class_, _ = import_context
    content = b"student_code,full_name\nA1,First\nA2,Second\n"
    digest = upload(client, content, class_id=class_.id).json()["digest"]

    def fail_second(_mapper, _connection, target):
        if target.student_code == "A2":
            raise IntegrityError("simulated conflict", {}, Exception("conflict"))

    event.listen(Student, "before_insert", fail_second)
    try:
        response = upload(client, content, class_id=class_.id, dry_run="false", expected_digest=digest)
        assert response.status_code == 409
        assert db.scalar(select(Student)) is None
        assert db.scalar(select(ClassEnrollment)) is None
        assert db.scalar(select(AuditLog)) is None
    finally:
        event.remove(Student, "before_insert", fail_second)


def test_size_row_limits_and_invalid_excel_are_rejected(monkeypatch):
    monkeypatch.setattr("backend.app.services.student_import.MAX_ROWS", 2)
    with pytest.raises(HTTPException):
        parse_students("data.csv", b"student_code,full_name\nA,A\nB,B\nC,C\n")
    for filename in ("data.xls", "data.xlsx"):
        with pytest.raises(HTTPException):
            parse_students(filename, b"not an excel file")
    monkeypatch.setattr("backend.app.services.student_import.MAX_BYTES", 10)
    with pytest.raises(HTTPException) as error:
        parse_students("data.csv", b"student_code,full_name\n")
    assert error.value.status_code == 413
